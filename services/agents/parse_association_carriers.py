"""Parse carrier contacts from KazATO (KZ) and BAMAP (BY) association pages.

Usage:
  python -m agents.parse_association_carriers
  python -m agents.parse_association_carriers --dry-run
  python -m agents.parse_association_carriers --kazato-only
  python -m agents.parse_association_carriers --bamap-only --enrich-emails
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
import time
import unicodedata
from datetime import date
from pathlib import Path
from typing import Any
from urllib.parse import unquote, urljoin, urlparse

import httpx
from bs4 import BeautifulSoup
from openpyxl import load_workbook

try:
    import xlrd  # legacy .xls (KazATO currently publishes .xls)
except ImportError:  # pragma: no cover
    xlrd = None  # type: ignore

SERVICES_ROOT = Path(__file__).resolve().parents[1]
REPO_ROOT = Path(__file__).resolve().parents[2]
if str(SERVICES_ROOT) not in sys.path:
    sys.path.insert(0, str(SERVICES_ROOT))

KAZATO_PAGE = "https://www.kazato.kz/pages/cpisok-perevozchikov"
BAMAP_PAGE = "https://bamap.org/activities/about/chlen_b/"
DEFAULT_OUT = REPO_ROOT / "data" / "suppliers_associations_cis.json"
RAW_XLSX = REPO_ROOT / "data" / "raw" / "kazato_latest.xlsx"
RAW_XLS = REPO_ROOT / "data" / "raw" / "kazato_latest.xls"
_EXCEL_EXTS = (".xlsx", ".xls")
ENRICH_REPORT = REPO_ROOT / "data" / "enrich_emails_report.json"

EMAIL_RE = re.compile(r"[a-zA-Z0-9._%+-]+@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}")
EMAIL_SKIP = ("example.com", "sentry.io", "wixpress", "schema.org", "png", "jpg", "jpeg", "gif")
SECTION_SKIP = ("пассажирские перевозки",)
SECTION_INCLUDE = ("грузовые перевозки", "ассоциированные члены")

USER_AGENT = (
    "Mozilla/5.0 (compatible; AutoLogisticsOS/1.0; +https://transinvest.ru) "
    "association-carrier-parser"
)


def _slug(text: str, *, prefix: str = "") -> str:
    s = unicodedata.normalize("NFKD", text)
    ascii_s = s.encode("ascii", "ignore").decode("ascii")
    ascii_s = re.sub(r"[^a-zA-Z0-9]+", "_", ascii_s.lower()).strip("_")
    ascii_s = re.sub(r"_+", "_", ascii_s)
    if not ascii_s:
        ascii_s = hashlib.md5(text.encode("utf-8")).hexdigest()[:10]
    if prefix:
        return f"{prefix}_{ascii_s}"
    return ascii_s


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _extract_emails(*parts: str) -> list[str]:
    found: list[str] = []
    seen: set[str] = set()
    for part in parts:
        if not part:
            continue
        for match in EMAIL_RE.findall(part):
            email = match.lower().strip(".")
            domain = email.split("@")[-1]
            if any(skip in domain for skip in EMAIL_SKIP):
                continue
            if email not in seen:
                seen.add(email)
                found.append(email)
    return found


def _normalize_kz_phone(raw: str) -> str | None:
    digits = re.sub(r"\D", "", raw)
    if len(digits) < 10:
        return None
    if digits.startswith("8") and len(digits) == 11:
        digits = "7" + digits[1:]
    elif len(digits) == 10:
        digits = "7" + digits
    elif len(digits) > 11:
        digits = digits[-11:]
    if len(digits) != 11 or not digits.startswith("7"):
        return None
    return f"+{digits[0]} {digits[1:4]} {digits[4:7]} {digits[7:9]} {digits[9:11]}"


def _normalize_by_phone(raw: str) -> str | None:
    digits = re.sub(r"\D", "", raw)
    if len(digits) < 9:
        return None
    if digits.startswith("375"):
        body = digits[3:]
    else:
        body = digits[-9:]
    if len(body) != 9:
        return None
    return f"+375 {body[0:2]} {body[2:5]} {body[5:7]} {body[7:9]}"


def _split_phones(raw: str, *, country: str) -> list[str]:
    if not raw:
        return []
    parts = re.split(r"[,;/]|(?<=\))\s+(?=\()", raw)
    norm_fn = _normalize_kz_phone if country == "kz" else _normalize_by_phone
    out: list[str] = []
    seen: set[str] = set()
    for part in parts:
        phone = norm_fn(part.strip())
        if phone and phone not in seen:
            seen.add(phone)
            out.append(phone)
    return out


def _city_from_address(address: str) -> str:
    m = re.search(r"г\.\s*([^,]+)", address, re.I)
    return m.group(1).strip() if m else ""


def _is_excel_href(href: str) -> bool:
    low = unquote(href or "").lower()
    return any(low.endswith(ext) or f"{ext}?" in low for ext in _EXCEL_EXTS)


def fetch_kazato_excel_url(client: httpx.Client) -> str:
    resp = client.get(KAZATO_PAGE)
    resp.raise_for_status()
    soup = BeautifulSoup(resp.text, "html.parser")
    candidates: list[str] = []
    for a in soup.find_all("a", href=True):
        href = a["href"]
        if not _is_excel_href(href):
            continue
        text = (a.get_text() or "").lower()
        href_l = unquote(href).lower()
        if "adr" in href_l:
            continue
        if "перевозчик" in text or "перевозчик" in href_l or "список" in text:
            candidates.append(urljoin(KAZATO_PAGE, href))
    if not candidates:
        for a in soup.find_all("a", href=True):
            href = a["href"]
            href_l = unquote(href).lower()
            if _is_excel_href(href) and "adr" not in href_l:
                candidates.append(urljoin(KAZATO_PAGE, href))
    if not candidates:
        raise RuntimeError("KazATO Excel link not found on page")
    # Prefer .xls/.xlsx whose URL mentions перевозчик
    ranked = sorted(
        candidates,
        key=lambda u: (
            0 if "перевозчик" in unquote(u).lower() else 1,
            0 if u.lower().endswith(".xls") or u.lower().endswith(".xlsx") else 1,
        ),
    )
    return ranked[0]


def fetch_kazato_excel(client: httpx.Client, *, cache_path: Path | None = None) -> bytes:
    url = fetch_kazato_excel_url(client)
    resp = client.get(url)
    resp.raise_for_status()
    data = resp.content
    if cache_path is None:
        return data
    # Persist with the real extension (.xls vs .xlsx).
    suffix = ".xls" if url.lower().endswith(".xls") or data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" else ".xlsx"
    out = cache_path if cache_path.suffix.lower() == suffix else cache_path.with_suffix(suffix)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(data)
    # Keep legacy xlsx path in sync when we still got xlsx.
    if suffix == ".xlsx" and out != RAW_XLSX:
        RAW_XLSX.write_bytes(data)
    if suffix == ".xls":
        RAW_XLS.write_bytes(data)
    return data


def _workbook_is_xls(workbook_bytes: bytes) -> bool:
    return workbook_bytes[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def _sheet_names(workbook_bytes: bytes) -> list[str]:
    if _workbook_is_xls(workbook_bytes):
        if xlrd is None:
            raise RuntimeError("xlrd is required to read KazATO .xls files")
        return list(xlrd.open_workbook(file_contents=workbook_bytes).sheet_names())
    wb = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    return list(wb.sheetnames)


def _sheet_rows(workbook_bytes: bytes, sheet_name: str) -> list[dict[str, str]]:
    if _workbook_is_xls(workbook_bytes):
        if xlrd is None:
            raise RuntimeError("xlrd is required to read KazATO .xls files")
        book = xlrd.open_workbook(file_contents=workbook_bytes)
        names = book.sheet_names()
        if sheet_name not in names:
            target = sheet_name.strip().lower()
            sheet_name = next((n for n in names if n.strip().lower() == target), sheet_name)
        sh = book.sheet_by_name(sheet_name)
        if sh.nrows < 1:
            return []
        headers = [_clean_cell(sh.cell_value(0, c)) for c in range(sh.ncols)]
        out: list[dict[str, str]] = []
        for r in range(1, sh.nrows):
            item = {
                headers[c]: _clean_cell(sh.cell_value(r, c))
                for c in range(min(len(headers), sh.ncols))
            }
            if any(item.values()):
                out.append(item)
        return out

    wb = load_workbook(io.BytesIO(workbook_bytes), read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        # fuzzy match by stripped name
        target = sheet_name.strip().lower()
        sheet_name = next((n for n in wb.sheetnames if n.strip().lower() == target), sheet_name)
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [_clean_cell(h) for h in rows[0]]
    out = []
    for row in rows[1:]:
        item = {headers[i]: _clean_cell(row[i]) for i in range(min(len(headers), len(row)))}
        if any(item.values()):
            out.append(item)
    return out


def parse_kazato_workbook(workbook_bytes: bytes) -> list[dict[str, Any]]:
    names = _sheet_names(workbook_bytes)
    carrier_sheet = next((n for n in names if "перевоз" in n.lower()), names[0])
    reefer_sheet = next((n for n in names if "реф" in n.lower()), None)

    carriers = _sheet_rows(workbook_bytes, carrier_sheet)
    reefers = _sheet_rows(workbook_bytes, reefer_sheet) if reefer_sheet else []
    reefer_by_id = {
        r.get("Уникальный номер перевозчика", ""): r for r in reefers if r.get("Уникальный номер перевозчика")
    }

    suppliers: list[dict[str, Any]] = []
    for row in carriers:
        name = row.get("Перевозчик", "")
        carrier_id = row.get("Уникальный номер перевозчика", "")
        if not name:
            continue
        city = row.get("Город", "")
        address = row.get("Адрес офиса (строка 1)", "")
        zip_code = row.get("Почтовый индекс", "")
        phone_raw = row.get("Телефон #", "")
        email_raw = row.get("Электронный адрес", "")
        fax_raw = row.get("Факс #", "")

        reefer_row = reefer_by_id.get(carrier_id, {})
        reefer_count = reefer_row.get("Рефрижераторы", "")
        if reefer_row:
            phone_raw = phone_raw or reefer_row.get("Телефон #", "")
            email_raw = email_raw or reefer_row.get("Электронный адрес", "")
            fax_raw = fax_raw or reefer_row.get("Факс #", "")

        emails = _extract_emails(email_raw, fax_raw)
        phones = _split_phones(phone_raw, country="kz")
        notes = ["КазАТО"]
        if city:
            notes.append(city)
        if address:
            notes.append(address)
        if zip_code:
            notes.append(f"index:{zip_code}")
        if reefer_count:
            notes.append(f"reefer_count:{reefer_count}")

        code_suffix = _slug(carrier_id.replace("/", "_")) if carrier_id else _slug(name)
        suppliers.append(
            {
                "name": name,
                "code": f"kazato_{code_suffix}",
                "modes": ["ftl_truck"],
                "corridors": ["international"],
                "phones": phones,
                "emails": emails,
                "contacts": [],
                "notes": notes,
                "source": "kazato_2026",
                "carrier_id": carrier_id,
            }
        )
    return suppliers


def _parse_bamap_row(cells: list[str]) -> dict[str, Any] | None:
    if len(cells) < 4:
        return None
    num_raw, name, address, phone_raw = cells[0], cells[1], cells[2], cells[3]
    if not name or not re.match(r"^\d+$", num_raw):
        return None
    city = _city_from_address(address)
    phones = _split_phones(phone_raw, country="by")
    notes = ["БАМАП"]
    if address:
        notes.append(address)
    if city:
        notes.append(city)
    return {
        "name": name,
        "code": f"bamap_{num_raw}_{_slug(name)}",
        "modes": ["ftl_truck"],
        "corridors": ["international"],
        "phones": phones,
        "emails": [],
        "contacts": [],
        "notes": notes,
        "source": "bamap_2026",
        "city": city,
    }


def parse_bamap_members(html: str) -> list[dict[str, Any]]:
    soup = BeautifulSoup(html, "html.parser")
    table = soup.find("table")
    if not table:
        raise RuntimeError("BAMAP member table not found")

    section = ""
    suppliers: list[dict[str, Any]] = []
    for tr in table.find_all("tr"):
        cells = [c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"])]
        if not cells:
            continue
        joined = " ".join(cells).lower()
        if len(cells) <= 2 and any(k in joined for k in SECTION_INCLUDE + SECTION_SKIP):
            if any(k in joined for k in SECTION_SKIP):
                section = "skip"
            elif "ассоциирован" in joined:
                section = "associated"
            elif "грузов" in joined:
                section = "freight"
            continue
        if section == "skip":
            continue
        if section not in ("freight", "associated"):
            continue
        record = _parse_bamap_row(cells)
        if record:
            if section == "associated":
                record["notes"].insert(1, "ассоциированный член")
            suppliers.append(record)
    return suppliers


def _search_urls(client: httpx.Client, query: str, *, limit: int = 3) -> list[str]:
    resp = client.post(
        "https://html.duckduckgo.com/html/",
        data={"q": query, "kl": "ru-ru"},
        headers={"Content-Type": "application/x-www-form-urlencoded"},
    )
    if resp.status_code >= 400:
        return []
    soup = BeautifulSoup(resp.text, "html.parser")
    urls: list[str] = []
    for a in soup.select("a.result__a"):
        href = a.get("href", "")
        if not href.startswith("http"):
            continue
        host = urlparse(href).netloc.lower()
        if "duckduckgo.com" in host:
            continue
        if href not in urls:
            urls.append(href)
        if len(urls) >= limit:
            break
    return urls


def _scrape_emails_from_url(client: httpx.Client, url: str) -> list[str]:
    try:
        resp = client.get(url)
        resp.raise_for_status()
    except httpx.HTTPError:
        return []
    emails = _extract_emails(resp.text)
    if emails:
        return emails
    parsed = urlparse(url)
    base = f"{parsed.scheme}://{parsed.netloc}"
    for suffix in ("/contacts", "/contact", "/kontakty", "/about", "/o-kompanii"):
        try:
            resp = client.get(base + suffix)
            if resp.status_code >= 400:
                continue
            emails = _extract_emails(resp.text)
            if emails:
                return emails
        except httpx.HTTPError:
            continue
    return []


def enrich_emails_from_web(
    suppliers: list[dict[str, Any]],
    *,
    delay_sec: float = 2.0,
    limit: int = 0,
) -> tuple[list[dict[str, Any]], dict[str, Any]]:
    report: dict[str, Any] = {"found": [], "not_found": [], "errors": []}
    enriched = [dict(s) for s in suppliers]
    targets = [s for s in enriched if not s.get("emails")]
    if limit > 0:
        targets = targets[:limit]

    with httpx.Client(
        timeout=8.0,
        follow_redirects=True,
        headers={"User-Agent": USER_AGENT},
    ) as client:
        for supplier in targets:
            name = supplier.get("name", "")
            city = supplier.get("city") or next(
                (n for n in supplier.get("notes", []) if n and not n.startswith("index:")), ""
            )
            query = f'"{name}" {city} официальный сайт'.strip()
            entry = {"code": supplier.get("code"), "name": name, "query": query}
            try:
                urls = _search_urls(client, query, limit=2)
                emails: list[str] = []
                for url in urls:
                    emails = _scrape_emails_from_url(client, url)
                    if emails:
                        entry["url"] = url
                        break
                if emails:
                    supplier["emails"] = emails
                    supplier.setdefault("notes", []).append(f"email_source:{entry.get('url', '')}")
                    entry["emails"] = emails
                    report["found"].append(entry)
                else:
                    report["not_found"].append(entry)
            except Exception as exc:  # noqa: BLE001
                entry["error"] = str(exc)
                report["errors"].append(entry)
            time.sleep(delay_sec)
    return enriched, report


def build_catalog(
    *,
    kazato: bool = True,
    bamap: bool = True,
    enrich_emails: bool = False,
    enrich_limit: int = 0,
    cache_xlsx: bool = True,
) -> tuple[dict[str, Any], dict[str, Any] | None]:
    suppliers: list[dict[str, Any]] = []
    enrich_report: dict[str, Any] | None = None

    with httpx.Client(
        timeout=30.0,
        follow_redirects=True,
        headers={"User-Agent": USER_AGENT},
    ) as client:
        if kazato:
            xlsx = fetch_kazato_excel(client, cache_path=RAW_XLSX if cache_xlsx else None)
            suppliers.extend(parse_kazato_workbook(xlsx))
        if bamap:
            resp = client.get(BAMAP_PAGE)
            resp.raise_for_status()
            suppliers.extend(parse_bamap_members(resp.text))

    if bamap and enrich_emails:
        by_only = [s for s in suppliers if s.get("source") == "bamap_2026"]
        others = [s for s in suppliers if s.get("source") != "bamap_2026"]
        enriched_by, enrich_report = enrich_emails_from_web(by_only, limit=enrich_limit)
        suppliers = others + enriched_by

    with_email = sum(1 for s in suppliers if s.get("emails"))
    for s in suppliers:
        s.pop("carrier_id", None)
        s.pop("city", None)
    catalog: dict[str, Any] = {
        "source": "associations_kz_by",
        "imported_at": date.today().isoformat(),
        "counts": {
            "total": len(suppliers),
            "with_email": with_email,
            "kazato": sum(1 for s in suppliers if s.get("source") == "kazato_2026"),
            "bamap": sum(1 for s in suppliers if s.get("source") == "bamap_2026"),
        },
        "suppliers": suppliers,
    }
    return catalog, enrich_report


def main() -> None:
    parser = argparse.ArgumentParser(description="Parse KazATO and BAMAP carrier lists")
    parser.add_argument("--dry-run", action="store_true", help="Print stats only, do not write JSON")
    parser.add_argument("--kazato-only", action="store_true")
    parser.add_argument("--bamap-only", action="store_true")
    parser.add_argument(
        "--enrich-emails",
        action="store_true",
        help="Search company websites for BY emails (default when --bamap-only or full run)",
    )
    parser.add_argument("--no-enrich-emails", action="store_true")
    parser.add_argument("--enrich-limit", type=int, default=0, help="Limit BY email enrichment count")
    parser.add_argument("--out", type=Path, default=DEFAULT_OUT)
    args = parser.parse_args()

    kazato = not args.bamap_only
    bamap = not args.kazato_only
    enrich = args.enrich_emails or (bamap and not args.no_enrich_emails and not args.dry_run)

    catalog, enrich_report = build_catalog(
        kazato=kazato,
        bamap=bamap,
        enrich_emails=enrich,
        enrich_limit=args.enrich_limit,
        cache_xlsx=not args.dry_run,
    )

    summary = {
        "counts": catalog["counts"],
        "out": str(args.out),
        "enrich_emails": enrich,
        "sample": catalog["suppliers"][:3],
    }

    if args.dry_run:
        print(json.dumps(summary, ensure_ascii=False, indent=2))
        return

    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps(catalog, ensure_ascii=False, indent=2), encoding="utf-8")
    if enrich_report is not None:
        ENRICH_REPORT.parent.mkdir(parents=True, exist_ok=True)
        ENRICH_REPORT.write_text(json.dumps(enrich_report, ensure_ascii=False, indent=2), encoding="utf-8")
        summary["enrich_report"] = str(ENRICH_REPORT)
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
